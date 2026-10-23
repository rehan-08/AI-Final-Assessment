"""
Builds report/Final_Report.docx from the tables in outputs/ and figures in visualization/.
Run after the notebook has been executed:  python src/build_report.py
"""
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT, VIS = ROOT / "outputs", ROOT / "visualization"
(ROOT / "report").mkdir(exist_ok=True)

from observations import REPORT  # narrative text kept next to the notebook observations

lab = pd.read_csv(OUT / "labeled_messages.csv", parse_dates=["date"])
scores = pd.read_csv(OUT / "monthly_scores.csv")
top_pos = pd.read_csv(OUT / "top3_positive_by_month.csv")
top_neg = pd.read_csv(OUT / "top3_negative_by_month.csv")
overall = pd.read_csv(OUT / "overall_employee_ranking.csv")
risk = pd.read_csv(OUT / "flight_risk.csv")
metrics = pd.read_csv(OUT / "model_metrics.csv", index_col=0)
coefs = pd.read_csv(OUT / "model_coefficients.csv", index_col=0)

doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)


def para(text, bold=False, italic=False):
    p = doc.add_paragraph(); r = p.add_run(text); r.bold = bold; r.italic = italic
    return p


def bullets(items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def table(df, col_widths=None, font=9):
    t = doc.add_table(rows=1, cols=len(df.columns)); t.style = "Light Grid Accent 1"
    for i, c in enumerate(df.columns):
        t.rows[0].cells[i].text = str(c)
    for _, row in df.iterrows():
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = f"{v:.3f}" if isinstance(v, float) else str(v)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(font)
    doc.add_paragraph()


def figure(name, caption, width=6.3):
    doc.add_picture(str(VIS / name), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = para(caption, italic=True); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.runs[0].font.size = Pt(9); c.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)


short = lambda e: e.split("@")[0]

# ----------------------------------------------------------------------------- title
doc.add_heading("Employee Sentiment Analysis – Final Report", 0)
para(REPORT["subtitle"], italic=True)

doc.add_heading("Executive summary", 1)
bullets(REPORT["exec"])

# ----------------------------------------------------------------------------- 1
doc.add_heading("1. Data and approach", 1)
para(REPORT["data"])
doc.add_heading("1.1 Sentiment labeling methodology (Task 1)", 2)
for p in REPORT["labeling"]:
    para(p)
agree = pd.read_csv(OUT / "labeler_agreement.csv")
table(agree)
figure("01_label_validation.png", "Figure 1 – Transformer labels vs VADER and TextBlob, and transformer confidence per label.")

# ----------------------------------------------------------------------------- 2
doc.add_heading("2. Exploratory data analysis (Task 2)", 1)
bullets(REPORT["eda"])
figure("02_sentiment_distribution.png", "Figure 2 – Overall sentiment distribution and sentiment mix per employee.")
figure("03_sentiment_trend_monthly.png", "Figure 3 – Monthly volume by sentiment and monthly net sentiment.")
figure("04_weekday_year_patterns.png", "Figure 4 – Sentiment share by weekday and by year.")
figure("05_length_and_heatmap.png", "Figure 5 – Message length by sentiment; net sentiment per message by employee and quarter.")

# ----------------------------------------------------------------------------- 3
doc.add_heading("3. Employee scoring (Task 3)", 1)
bullets(REPORT["scoring"])
figure("06_monthly_score_heatmap.png", "Figure 6 – Monthly sentiment score for every employee (resets each month).", 6.5)

# ----------------------------------------------------------------------------- 4
doc.add_heading("4. Employee ranking (Task 4)", 1)
bullets(REPORT["ranking"])
doc.add_heading("4.1 Monthly top 3 (all 24 months)", 2)
wide = lambda t: (t.assign(cell=t.employee.map(short) + " (" + t.score.astype(str) + ")")
                   .pivot(index="month", columns="rank", values="cell"))
side = pd.concat([wide(top_pos).add_prefix("Pos #"), wide(top_neg).add_prefix("Neg #")], axis=1).reset_index()
side.columns = ["Month", "Positive #1", "Positive #2", "Positive #3", "Negative #1", "Negative #2", "Negative #3"]
table(side, font=7.5)
doc.add_heading("4.2 Overall ranking across the period", 2)
o = overall.copy(); o["employee"] = o.employee.map(short); o["avg_monthly_score"] = o.avg_monthly_score.round(2)
table(o)
figure("07_employee_rankings.png", "Figure 7 – Total score and number of months in a top-3 list.")

# ----------------------------------------------------------------------------- 5
doc.add_heading("5. Flight-risk identification (Task 5)", 1)
bullets(REPORT["flight"])
r = risk.copy(); r["employee"] = r.employee.map(short)
for c in ["peak_window_start", "peak_window_end", "first_flagged_on"]:
    r[c] = pd.to_datetime(r[c]).dt.strftime("%Y-%m-%d").fillna("-")
table(r)
figure("08_flight_risk_rolling.png", "Figure 8 – Rolling 30-day count of negative messages; dashed line = threshold of 4.")
figure("09_flight_risk_peak.png", "Figure 9 – Peak negatives in any 30-day window per employee.", 5.5)

# ----------------------------------------------------------------------------- 6
doc.add_heading("6. Predictive model (Task 6)", 1)
bullets(REPORT["model_setup"])
doc.add_heading("6.1 Evaluation", 2)
table(metrics.reset_index().rename(columns={"index": "Evaluation"}))
doc.add_heading("6.2 Coefficients (standardized features, OLS on all rows)", 2)
table(coefs.reset_index().rename(columns={"index": "feature"}))
bullets(REPORT["model_findings"])
figure("10_feature_correlation.png", "Figure 10 – Correlation of features with the monthly score.", 5.2)
figure("11_model_performance.png", "Figure 11 – Actual vs predicted, residuals and standardized coefficients.")
figure("12_model_metrics.png", "Figure 12 – Evaluation metrics across splits.", 5.2)

# ----------------------------------------------------------------------------- 7
doc.add_heading("7. Recommendations", 1)
bullets(REPORT["recommendations"])
doc.add_heading("8. Assumptions and limitations", 1)
bullets(REPORT["limitations"])
doc.add_heading("9. Reproducibility", 1)
para(REPORT["repro"])

doc.save(ROOT / "report" / "Final_Report.docx")
print("report written")
